"""
Analizador de tesis doctoral usando Gemini API.
Escribe resultados directamente al .docx de la tesis.
"""

import os, sys, time, shutil, zipfile, xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

# Fix: asegurar lxml de usuario antes que el del sistema
sys.path.insert(0, str(Path.home() / ".local/lib/python3.10/site-packages"))

import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from google import genai
from google.genai import errors as genai_errors
from dotenv import load_dotenv

BASE_DIR = Path(__file__).parent
load_dotenv(BASE_DIR / ".env")

API_KEY = os.getenv("GEMINI_API_KEY")
if not API_KEY or API_KEY == "tu_clave_aqui":
    print("ERROR: Agrega tu clave en el archivo .env")
    sys.exit(1)

client = genai.Client(api_key=API_KEY)
MODELOS = ["gemini-2.5-flash", "gemini-2.0-flash-lite", "gemini-2.5-pro"]
TESIS_PATH = BASE_DIR / "TESIS DOCTORAL 2025 SANTANA-RAMIREZ.docx"

# ─── Contexto fijo de la tesis ────────────────────────────────────────────────
CONTEXTO_TESIS = """
DATOS DE LA TESIS:
- Autor: José de Jesús Santana Ramírez (UAQ, Doctor en Ciencias de la Computación)
- Título: Modelo predictivo de histéresis de un sistema mecánico implementando
  técnicas de razonamiento basado en casos, para sistemas dinámicos probabilísticos discretos
- Director: Dr. Juan Carlos Antonio Jáuregui Correa
- Codirector: Dr. Fausto Abraham Jacques García

HIPÓTESIS (3.2): La integración de CBR + aritmética modular + análisis de entropía,
aplicado a datos experimentales, reducirá los errores de predicción de histéresis
en sistemas mecánicos bajo condiciones variables e incertidumbre.

MODELOS FÍSICOS: LuGre, Bouc-Wen, Dahl, Preisach, Prandtl-Ishlinskii
TÉCNICAS IA: KAN, HiPPO-KAN, CBR, PINNs (comparativas)
HARDWARE: NI cDAQ-9174 (NI 9205, NI 9234, NI 9219), fresadora CNC, shaker TIRA

RESULTADOS CLAVE (Dic 2025):
- Correlación F vs aceleración directa: r=0.004 (nula)
- Correlación F vs envolvente de aceleración: r=0.822 (alta)
- Modelo lineal envolvente: R²=0.675; histéresis no-lineal: 32.5% varianza restante
- Sistema dinámico: fn=45.1 Hz, ζ=0.159, m=12.42 kg, k=997,548 N/m, Fc=2.55 N
- Arquitectura KAN-PINN: F = m·a + k·x + c·v + α·k·E + (1-α)·k·z
- Artículo IFToMM: "Detección de histéresis F-v a partir de vibraciones"
"""


# ─── I/O de documentos ────────────────────────────────────────────────────────
def leer_docx(path: str) -> str:
    with zipfile.ZipFile(path) as z:
        with z.open("word/document.xml") as f:
            tree = ET.parse(f)
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    parrafos = []
    for p in tree.getroot().iter(f"{{{W}}}p"):
        texto = "".join(t.text for t in p.iter(f"{{{W}}}t") if t.text).strip()
        if texto:
            parrafos.append(texto)
    return "\n".join(parrafos)


def leer(path: str) -> str:
    with open(path, encoding="utf-8", errors="ignore") as f:
        return f.read()


def cargar_contexto() -> str:
    tesis = leer_docx(str(TESIS_PATH))
    cap4  = leer(str(BASE_DIR / "Capitulo4_Resultados_Extendido.md"))
    return f"{CONTEXTO_TESIS}\n\n=== TESIS COMPLETA ===\n{tesis[:25000]}\n\n=== CAP4 ===\n{cap4[:8000]}"


# ─── Escritura directa al .docx ───────────────────────────────────────────────
SECCIONES = {
    "resumen":      "Resumen",
    "abstract":     "Abstract",
    "cap1":         "Capítulo 1",
    "cap3_preg":    "3.1",
    "cap3_prop":    "3.2",
    "cap3_obj":     "3.3",
    "cap4":         "Capítulo 4",
    "cap5":         "Capítulo 5",
    "cap6":         "Capítulo 6",
}


def backup_tesis():
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = BASE_DIR / f"TESIS_backup_{ts}.docx"
    shutil.copy2(TESIS_PATH, dest)
    print(f"  Backup creado: {dest.name}")
    return dest


def escribir_en_seccion(seccion_key: str, nuevo_texto: str):
    """Reemplaza el contenido de una sección en el .docx de la tesis."""
    if seccion_key not in SECCIONES:
        print(f"Sección desconocida: {seccion_key}")
        return

    encabezado = SECCIONES[seccion_key]
    backup_tesis()

    doc = docx.Document(str(TESIS_PATH))
    parrafos = doc.paragraphs
    idx_inicio = None

    # Buscar el encabezado de la sección
    for i, p in enumerate(parrafos):
        if encabezado.lower() in p.text.lower():
            idx_inicio = i
            break

    if idx_inicio is None:
        print(f"  No se encontró la sección '{encabezado}' en la tesis.")
        return

    # Borrar párrafos de la sección hasta el siguiente encabezado
    i = idx_inicio + 1
    while i < len(parrafos):
        p = parrafos[i]
        estilo = p.style.name.lower()
        if ("heading" in estilo or "título" in estilo or
                any(p.text.strip().startswith(h) for h in
                    ["Capítulo", "Abstract", "Resumen", "1.", "2.", "3.", "4.", "5.", "6."])):
            break
        # Limpiar párrafo
        for run in p.runs:
            run.text = ""
        i += 1

    # Insertar nuevo contenido después del encabezado
    p_ref = parrafos[idx_inicio]
    for linea in nuevo_texto.strip().split("\n"):
        if not linea.strip():
            continue
        nuevo_p = docx.oxml.OxmlElement('w:p')
        p_ref._element.addnext(nuevo_p)
        new_para = docx.text.paragraph.Paragraph(nuevo_p, doc)
        run = new_para.add_run(linea.strip())
        run.font.size = Pt(12)
        new_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.save(str(TESIS_PATH))
    print(f"  ✓ Sección '{encabezado}' actualizada en la tesis.")


def menu_escritura(resultado: str):
    """Pregunta al usuario si quiere escribir el resultado al .docx."""
    print("\n¿Dónde escribir este resultado en la tesis?")
    for k, v in SECCIONES.items():
        print(f"  [{k:12s}] → {v}")
    print("  [no          ] → No escribir")
    opcion = input("\nSección: ").strip().lower()
    if opcion != "no" and opcion in SECCIONES:
        escribir_en_seccion(opcion, resultado)
    elif opcion != "no":
        nombre = f"resultado_{opcion}_{datetime.now().strftime('%H%M%S')}.txt"
        with open(BASE_DIR / nombre, "w", encoding="utf-8") as f:
            f.write(resultado)
        print(f"  Guardado en: {nombre}")


# ─── Motor Gemini ─────────────────────────────────────────────────────────────
def gemini(pregunta: str, contexto: str, titulo: str = "") -> str:
    if titulo:
        print(f"\n{'='*60}\n  {titulo}\n{'='*60}")

    prompt = f"""Eres un experto en mecatrónica, control no lineal y machine learning científico.
Analiza el siguiente material de tesis doctoral sobre predicción de histéresis.

{contexto[:32000]}

TAREA:
{pregunta}

Responde en español. Sé estructurado, concreto y accionable.
"""
    for modelo in MODELOS:
        for intento in range(3):
            try:
                print(f"  [modelo: {modelo}]")
                r = client.models.generate_content(model=modelo, contents=prompt)
                print(r.text)
                return r.text
            except genai_errors.ServerError:
                espera = 20 * (intento + 1)
                print(f"  Servidor ocupado, reintentando en {espera}s...")
                time.sleep(espera)
            except genai_errors.ClientError as e:
                if "429" in str(e):
                    print(f"  Cuota agotada en {modelo}, probando siguiente...")
                    break
                raise
    return ""


# ─── Análisis predefinidos ────────────────────────────────────────────────────
ANALISIS = {
    "1": (
        "Identifica las principales fortalezas y brechas metodológicas de esta tesis. "
        "¿Qué falta para que sea publicable en una revista Q1? Lista máximo 5 puntos críticos "
        "con acciones concretas para cada uno.",
        "ANÁLISIS COMPLETO — BRECHAS Y FORTALEZAS"
    ),
    "2": (
        "Evalúa la coherencia interna del Capítulo 3 (metodología). "
        "¿Los objetivos específicos están alineados con la hipótesis y el diseño experimental? "
        "¿Hay pasos metodológicos sin respaldo o mal definidos?",
        "COHERENCIA DEL CAPÍTULO 3 — METODOLOGÍA"
    ),
    "3": (
        "Analiza los resultados del Capítulo 4. "
        "¿Son suficientes para sustentar la hipótesis 3.2 (CBR+aritmética modular+entropía)? "
        "¿Qué experimentos adicionales se necesitan? Sé específico con métricas.",
        "EVALUACIÓN DEL CAPÍTULO 4 — RESULTADOS"
    ),
    "4": (
        """Diseña un DOE (Design of Experiments) completo para validar que el índice de histéresis
puede detectar el desgaste del cortador en fresado CNC. El experimento debe:

1. VARIABLES DE CONTROL (factores): Define niveles para:
   - Estado de desgaste de la herramienta (VB flank wear: nuevo, desgaste medio, desgastado)
   - Velocidad de corte (RPM)
   - Avance por diente (mm/tooth)
   - Profundidad de corte (mm)

2. VARIABLES DE RESPUESTA: Lista las métricas que medirá el modelo:
   - Índice de histéresis H = E_loop / (W_loop × ptp(F))
   - Correlación r(F, envolvente)
   - Parámetros Bouc-Wen identificados (α, A, β, γ)
   - Error de predicción del modelo CBR

3. DISEÑO EXPERIMENTAL: Recomienda el tipo de diseño (factorial completo, fraccionado,
   superficie de respuesta) justificando la elección según el número de corridas posibles
   con el hardware disponible (fresadora CNC, NI cDAQ-9174).

4. PROTOCOLO DE MEDICIÓN: Pasos concretos para capturar datos con el sistema actual:
   - Cuándo medir (inicio, mitad, fin de vida de la herramienta)
   - Qué señales adquirir (fuerza, aceleración, canales)
   - Cómo calcular el desgaste de referencia (microscopio óptico, VB)

5. ANÁLISIS ESTADÍSTICO: ANOVA, efectos principales, interacciones relevantes.

6. HIPÓTESIS ESPECÍFICA DEL DOE: Formula una hipótesis testable como:
   "El índice H aumenta X% por cada 0.1mm de desgaste VB bajo condiciones Y"

Basa todo en el hardware que ya tenemos: NI cDAQ-9174, módulos 9205/9234,
acelerómetro PCB 352C33, celda DYMH-105, fresadora CNC 2 dientes a 3720 RPM.""",
        "DOE — VALIDACIÓN DESGASTE DE CORTADOR"
    ),
    "5": (
        "Genera 10 preguntas difíciles que el comité doctoral podría hacer en la defensa. "
        "Incluye: fundamentos teóricos (CBR, aritmética modular, entropía), "
        "validación experimental, comparación con estado del arte, "
        "limitaciones y trabajo futuro. Para cada pregunta sugiere la respuesta clave.",
        "PREGUNTAS DE DEFENSA DOCTORAL"
    ),
}


# ─── Menú principal ───────────────────────────────────────────────────────────
def menu():
    print("\n" + "="*60)
    print("  ANALIZADOR DE TESIS — Gemini + Escritura Directa")
    print("="*60)
    print("[1] Análisis completo (brechas y fortalezas)")
    print("[2] Coherencia del Capítulo 3 (metodología)")
    print("[3] Evaluar resultados del Capítulo 4")
    print("[4] Diseñar DOE — validación desgaste del cortador")
    print("[5] Preguntas de defensa doctoral")
    print("[6] Redactar Resumen + Abstract (escribe al .docx)")
    print("[7] Pregunta libre → escribe al .docx")
    print("[0] Salir")
    return input("\nOpción: ").strip()


def main():
    print("\nCargando documentos de la tesis...")
    contexto = cargar_contexto()
    print(f"Contexto cargado: {len(contexto):,} caracteres\n")

    while True:
        opcion = menu()

        if opcion == "0":
            print("Hasta luego.")
            break

        elif opcion in ANALISIS:
            pregunta, titulo = ANALISIS[opcion]
            resultado = gemini(pregunta, contexto, titulo)
            if resultado:
                guardar = input("\n¿Escribir en la tesis o guardar? (s/n): ").strip().lower()
                if guardar == "s":
                    menu_escritura(resultado)

        elif opcion == "6":
            pregunta = (
                "Redacta el RESUMEN (español, 250 palabras) y el ABSTRACT (inglés, 250 palabras) "
                "para esta tesis. Estructura: contexto → problema → propuesta (CBR+aritmética "
                "modular+entropía) → metodología → resultados experimentales (r=0.822, fn=45.1 Hz, "
                "Bouc-Wen) → conclusión. Incluye palabras clave. "
                "NO menciones PINNs ni DeepONet como metodología principal."
            )
            resultado = gemini(pregunta, contexto, "NUEVO RESUMEN Y ABSTRACT")
            if resultado:
                print("\n¿Escribir el Resumen directamente en la tesis?")
                r = input("(s/n): ").strip().lower()
                if r == "s":
                    escribir_en_seccion("resumen", resultado)

        elif opcion == "7":
            pregunta = input("\nEscribe tu instrucción para Gemini: ").strip()
            if pregunta:
                resultado = gemini(pregunta, contexto, "CONSULTA LIBRE")
                if resultado:
                    guardar = input("\n¿Escribir en la tesis? (s/n): ").strip().lower()
                    if guardar == "s":
                        menu_escritura(resultado)

        else:
            print("Opción no válida.")


if __name__ == "__main__":
    main()
