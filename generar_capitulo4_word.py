"""
Generador de Capítulo 4 en formato Word con ecuaciones y tablas
Requiere: pip install python-docx
"""

from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_equation(paragraph, latex_code):
    """Agrega ecuación en formato LaTeX para conversión manual en Word"""
    run = paragraph.add_run(latex_code)
    run.font.name = 'Consolas'
    run.font.size = Pt(10)
    run.font.color.rgb = None  # Negro

def set_cell_shading(cell, color):
    """Aplica color de fondo a celda"""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading)

def create_table(doc, headers, rows, col_widths=None):
    """Crea tabla formateada"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Encabezados
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].text = header
        header_cells[i].paragraphs[0].runs[0].bold = True
        header_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(header_cells[i], 'D9E2F3')
    
    # Filas de datos
    for row_data in rows:
        row = table.add_row()
        for i, cell_text in enumerate(row_data):
            row.cells[i].text = str(cell_text)
            row.cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    return table

def main():
    doc = Document()
    
    # Configurar estilos
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    
    # ========== TÍTULO ==========
    title = doc.add_heading('Capítulo 4. Resultados', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # ========== 4.1 INTRODUCCIÓN ==========
    doc.add_heading('4.1 Introducción', level=1)
    doc.add_paragraph(
        'Los resultados presentados en este capítulo se alinean con los objetivos específicos '
        'planteados en la investigación. Se documentan los avances parciales obtenidos hasta la fecha, '
        'incluyendo el desarrollo de modelos matemáticos para histéresis, la implementación de técnicas '
        'de inteligencia artificial y la validación experimental del sistema de adquisición de datos.'
    )
    
    # ========== 4.2 PUBLICACIONES ==========
    doc.add_heading('4.2 Publicaciones y Difusión Científica', level=1)
    
    doc.add_heading('4.2.1 Artículo en Congreso IFToMM', level=2)
    doc.add_paragraph(
        'Se redactó y presentó el artículo titulado "Detección de histéresis F-v a partir de vibraciones: '
        'sensor virtual de fuerza y modelos Dahl/LuGre/Bouc-Wen en fresado de aluminio" en el congreso de IFToMM.'
    )
    
    doc.add_paragraph('Contribuciones principales del artículo:', style='List Bullet')
    
    p = doc.add_paragraph(style='List Number')
    p.add_run('Pipeline reproducible ').bold = True
    p.add_run('para detectar eventos en aceleración, reconstruir velocidad e inferir fuerza sintética '
              'mediante modelos de fricción con memoria (Dahl, LuGre, Bouc-Wen) con selección automática '
              'por índice de histéresis.')
    
    p = doc.add_paragraph(style='List Number')
    p.add_run('Cuantificación adimensional ').bold = True
    p.add_run('robusta a ganancias desconocidas, mediante área y ancho normalizados del lazo de histéresis:')
    
    # Ecuación 1
    eq1 = doc.add_paragraph()
    eq1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_equation(eq1, r'\tilde{E}_{\text{loop}} = \left|\oint \tilde{F} \, d\tilde{v}\right|, \quad \tilde{W}_{\text{loop}} = \max(\tilde{v}) - \min(\tilde{v})')
    
    doc.add_paragraph('Donde ~ denota normalización z-score.')
    
    p = doc.add_paragraph(style='List Number')
    p.add_run('Índice de histéresis (H-index) ').bold = True
    p.add_run('definido como:')
    
    # Ecuación 2
    eq2 = doc.add_paragraph()
    eq2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_equation(eq2, r'H = \frac{\tilde{E}_{\text{loop}}}{\tilde{W}_{\text{loop}} \cdot \text{ptp}(\tilde{F})}')
    
    p = doc.add_paragraph(style='List Number')
    p.add_run('Base metodológica ').bold = True
    p.add_run('para incorporar FRF/Kalman (tiempo real) y calibración absoluta de energía (J/ciclo).')
    
    doc.add_paragraph(
        'Estado actual: El artículo recibió retroalimentación detallada y se están implementando las '
        'correcciones sugeridas. Se está transcribiendo de LaTeX a Word para facilitar la revisión colaborativa.'
    )
    
    doc.add_heading('4.2.2 Reporte Técnico PINN', level=2)
    doc.add_paragraph(
        'Se completó el documento técnico "Modelado de levitador magnético con redes neuronales", '
        'que documenta la metodología y resultados del modelo PINN aplicado a un sistema de levitación '
        'magnética como banco de pruebas para la validación de la arquitectura propuesta.'
    )
    
    # ========== 4.3 SISTEMA EXPERIMENTAL ==========
    doc.add_heading('4.3 Sistema Experimental de Adquisición', level=1)
    
    doc.add_heading('4.3.1 Configuración del Hardware', level=2)
    doc.add_paragraph(
        'Se diseñó e implementó un sistema de adquisición de datos basado en el chasis NI cDAQ-9174 '
        'con los siguientes módulos:'
    )
    
    # Tabla de módulos
    create_table(doc, 
        ['Módulo', 'Función', 'Especificaciones'],
        [
            ['NI 9205', 'Fuerza (voltaje analógico)', '±10V, diferencial, ai0-ai3, 2.5 kHz'],
            ['NI 9234', 'Vibración (IEPE)', '4 canales, 51.2 kS/s, excitación 4mA'],
            ['NI 9219', 'Celda de carga (puente)', '24-bit, 100 S/s, excitación configurable']
        ]
    )
    doc.add_paragraph()
    
    doc.add_paragraph('Sensores utilizados:')
    doc.add_paragraph('Acelerómetros: PCB 352C33, sensibilidad 98.3-100 mV/g.', style='List Bullet')
    doc.add_paragraph('Celda de carga: DYMH-105 (500 kg, 1.7 mV/V) con acondicionador INA-4LC-8NTC (G=601).', style='List Bullet')
    doc.add_paragraph('Shaker: TIRA TV 51144IN con amplificador BAA 1000.', style='List Bullet')
    
    doc.add_heading('4.3.2 Interfaz de Caracterización Desarrollada', level=2)
    doc.add_paragraph(
        'Se desarrolló una interfaz gráfica completa (caracterizacion_sensor_fuerza.py, 1085 líneas) '
        'para la caracterización del sensor de fuerza con las siguientes funcionalidades:'
    )
    
    doc.add_paragraph('Características implementadas:', style='List Bullet')
    doc.add_paragraph('Adquisición simultánea de fuerza y aceleración (2 canales).', style='List Bullet')
    doc.add_paragraph('Visualización en tiempo real de señales, FFT y FRF.', style='List Bullet')
    doc.add_paragraph('Cálculo de transmisibilidad entre acelerómetros.', style='List Bullet')
    doc.add_paragraph('Validación F=ma integrada.', style='List Bullet')
    doc.add_paragraph('Diagramas de Lissajous para análisis de fase.', style='List Bullet')
    doc.add_paragraph('Guardado automático de experimentos con metadatos.', style='List Bullet')
    
    doc.add_paragraph('Resultados de validación de datos:')
    create_table(doc,
        ['Parámetro', 'Valor'],
        [
            ['Archivo', 'sesion_20250610_191323_fuerza_200_rpm_arana.csv'],
            ['Muestras', '41,900 por canal'],
            ['NaN/Inf', '0/0'],
            ['Tiempos monótonos', '✓'],
            ['Outliers (|z|>6)', '339 en ai0, 0 en ai1-ai3']
        ]
    )
    doc.add_paragraph()
    
    doc.add_heading('4.3.3 Análisis de Función de Respuesta en Frecuencia (FRF)', level=2)
    doc.add_paragraph('Se realizaron experimentos de caracterización a múltiples frecuencias de excitación:')
    
    create_table(doc,
        ['Frecuencia (Hz)', '|H(f)| (g/V)', 'Fase (°)', 'Coherencia'],
        [
            ['5', '12.80', '-114.1', '-'],
            ['10', '19.26', '158.5', '0.004'],
            ['20', '0.39', '75.9', '0.009'],
            ['40', '0.80', '-86.5', '-'],
            ['50', '8.38', '-29.2', '0.026'],
            ['70', '10.06', '173.9', '0.009'],
            ['80', '6.76', '25.4', '0.012'],
            ['120', '10.92', '-66.9', '0.047']
        ]
    )
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.add_run('Observación: ').bold = True
    p.add_run('La coherencia baja en algunas frecuencias indica la necesidad de mejorar la sincronización '
              'entre canales y aumentar el tiempo de adquisición para promediar más ciclos.')
    
    # ========== 4.4 MODELO PINN ==========
    doc.add_heading('4.4 Modelo Matemático de Histéresis con PINN', level=1)
    
    doc.add_heading('4.4.1 Arquitectura del Modelo', level=2)
    doc.add_paragraph(
        'Se implementó una Red Neuronal Informada por la Física (PINN) para modelar el comportamiento '
        'de histéresis. La arquitectura consta de:'
    )
    
    p = doc.add_paragraph()
    p.add_run('Modelo Base (MLP):').bold = True
    doc.add_paragraph('Entrada: tiempo normalizado t', style='List Bullet')
    doc.add_paragraph('Capas ocultas: 3 capas × 64 neuronas', style='List Bullet')
    doc.add_paragraph('Activación: tanh (diferenciable para AD)', style='List Bullet')
    doc.add_paragraph('Salida: posición ŷ, corriente î', style='List Bullet')
    
    p = doc.add_paragraph()
    p.add_run('Función de pérdida total:').bold = True
    
    eq = doc.add_paragraph()
    eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_equation(eq, r'\mathcal{L}_{\text{total}} = \mathcal{L}_{\text{data}} + \lambda \mathcal{L}_{\text{phys}}')
    
    doc.add_paragraph('Donde:')
    
    eq = doc.add_paragraph()
    eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_equation(eq, r'\mathcal{L}_{\text{data}} = \text{MSE}\left(\frac{\hat{y} - y}{\sigma_y}\right) + \text{MSE}\left(\frac{\hat{i} - i}{\sigma_i}\right)')
    
    eq = doc.add_paragraph()
    eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_equation(eq, r'\mathcal{L}_{\text{phys}} = \text{MSE}(R_{\text{mech}}) + \text{MSE}(R_{\text{elect}})')
    
    p = doc.add_paragraph()
    p.add_run('Residuales físicos:').bold = True
    
    doc.add_paragraph('Mecánico:', style='List Bullet')
    eq = doc.add_paragraph()
    eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_equation(eq, r'R_{\text{mech}} = m\ddot{y} - \left(\frac{1}{2}\frac{\partial L(y)}{\partial y}i^2 + mg\right)')
    
    doc.add_paragraph('Eléctrico:', style='List Bullet')
    eq = doc.add_paragraph()
    eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_equation(eq, r'R_{\text{elect}} = L(y)\dot{i} - \left(u - Ri - \frac{\partial L(y)}{\partial y}\dot{y}i\right)')
    
    doc.add_heading('4.4.2 Resultados Comparativos', level=2)
    doc.add_paragraph('Se evaluaron tres configuraciones del modelo:')
    
    create_table(doc,
        ['Modelo', 'Variable', 'RMSE', 'R²', 'Mejora vs Base'],
        [
            ['Base (solo datos)', 'Posición', '0.0057 m', '0.3432', '-'],
            ['', 'Corriente', '0.2282 A', '0.3283', '-'],
            ['PINN (λ=10⁻⁴)', 'Posición', '0.0049 m', '0.5105', '+48.7%'],
            ['', 'Corriente', '0.2010 A', '0.4788', '+45.8%'],
            ['PINN (λ=10⁻³)', 'Posición', '0.0049 m', '0.5058', '+47.4%'],
            ['', 'Corriente', '0.1968 A', '0.5003', '+52.4%']
        ]
    )
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.add_run('Hallazgo clave: ').bold = True
    p.add_run('La incorporación de restricciones físicas mejoró el R² de ~0.33 a ~0.50, demostrando '
              'que el modelo se beneficia significativamente de la información física.')
    
    doc.add_heading('4.4.3 Identificación de Parámetros Físicos', level=2)
    doc.add_paragraph('Se exploró la capacidad del PINN para identificar los parámetros del modelo de inductancia:')
    
    eq = doc.add_paragraph()
    eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_equation(eq, r'L(y) = k_0 + \frac{k}{1 + y/a}')
    
    create_table(doc,
        ['Parámetro', 'Valor Inicial', 'Valor Identificado', 'Cambio'],
        [
            ['k₀ (H)', '3.63×10⁻²', '9.92×10⁻³', '-72.7%'],
            ['k (H·m)', '3.50×10⁻³', '5.23×10⁻²', '+1394%'],
            ['a (m)', '5.20×10⁻³', '4.48×10⁻³', '-13.8%']
        ]
    )
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.add_run('Resultado importante: ').bold = True
    p.add_run('El modelo con parámetros entrenables mostró peor rendimiento predictivo (R² = 0.2983 vs 0.5058), '
              'sugiriendo convergencia a mínimos locales. Este hallazgo demuestra un desafío clave en PINNs: '
              'más grados de libertad no garantizan mejor rendimiento y pueden conducir a soluciones no físicas.')
    
    doc.add_heading('4.4.4 Arquitectura KAN-PINN Optimizada', level=2)
    doc.add_paragraph('Se desarrolló una versión eficiente basada en Kolmogorov-Arnold Networks (KAN):')
    
    p = doc.add_paragraph()
    p.add_run('Arquitectura:').bold = True
    doc.add_paragraph('Input (t) → KANLayer(1→64) → KANLayer(64→64) → KANLayer(64→64) → KANLayer(64→2) → Output (y, i)')
    
    p = doc.add_paragraph()
    p.add_run('Optimizaciones implementadas:').bold = True
    doc.add_paragraph('Diferencias finitas para derivadas de 2º orden (10x speedup).', style='List Bullet')
    doc.add_paragraph('Curriculum learning (transición gradual datos → física).', style='List Bullet')
    doc.add_paragraph('Adaptive weighting de λ_phys.', style='List Bullet')
    doc.add_paragraph('Clipping de parámetros físicos en límites razonables.', style='List Bullet')
    
    doc.add_paragraph('Tiempos de entrenamiento:')
    create_table(doc,
        ['GPU', 'Tiempo', 'Speedup'],
        [
            ['A100', '35-50 min', '10-15x'],
            ['T4 (gratis)', '1-1.5 h', '6-8x'],
            ['GTX 1060', '8+ h', '1x']
        ]
    )
    doc.add_paragraph()
    
    # ========== 4.5 DOE ==========
    doc.add_heading('4.5 Comparativa DOE: PINN Clásicos vs CNN', level=1)
    
    doc.add_heading('4.5.1 Diseño del Experimento', level=2)
    doc.add_paragraph('Se realizó un Diseño de Experimentos (DOE) comprehensivo comparando:')
    doc.add_paragraph('Modelos PINN clásicos: Bouc-Wen, LuGre, Dahl.', style='List Bullet')
    doc.add_paragraph('Modelos CNN: Arquitectura temporal optimizada para GPU.', style='List Bullet')
    doc.add_paragraph('Datos experimentales: Señales de fricción a 400 Hz de frecuencia de muestreo.')
    
    doc.add_heading('4.5.2 Resultados del DOE', level=2)
    
    create_table(doc,
        ['Enfoque', 'R² Promedio', 'Desv. Est.', 'Tasa Éxito', 'Tiempo Prom.'],
        [
            ['CNN DOE', '0.9716', '0.0757', '100% (11/11)', '37.09 s'],
            ['PINN Clásicos', '0.0256', '0.0541', '24.5%', '1879.4 s'],
            ['CNN Sesiones', '-0.0496', '-', '100%', '68.1 s']
        ]
    )
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.add_run('Mejor resultado CNN: ').bold = True
    p.add_run('R² = 0.999928')
    
    doc.add_heading('4.5.3 Análisis de Resultados', level=2)
    
    p = doc.add_paragraph()
    p.add_run('Cambio de paradigma demostrado:').bold = True
    doc.add_paragraph('Los enfoques data-driven (CNN) superan dramáticamente a modelos físicos clásicos para datos de alta frecuencia.', style='List Bullet')
    doc.add_paragraph('La aceleración por GPU permite un entrenamiento 50-100x más rápido.', style='List Bullet')
    doc.add_paragraph('La CNN captura patrones temporales complejos que los PINN clásicos no pueden modelar.', style='List Bullet')
    
    p = doc.add_paragraph()
    p.add_run('Implicaciones para tribología:').bold = True
    doc.add_paragraph('Datos experimentales modernos requieren enfoques modernos de ML.', style='List Bullet')
    doc.add_paragraph('Modelos clásicos (Bouc-Wen, LuGre, Dahl) son válidos solo para regímenes cuasi-estáticos.', style='List Bullet')
    doc.add_paragraph('CNN abre posibilidades para monitoreo de desgaste en tiempo real.', style='List Bullet')
    
    doc.add_heading('4.5.4 Optimizaciones GPU Implementadas', level=2)
    doc.add_paragraph('Para hardware con memoria limitada (3GB):')
    doc.add_paragraph('Modelo lightweight: 16→32→16 filtros (vs 64→128→64).', style='List Bullet')
    doc.add_paragraph('Datasets reducidos: 10k puntos máximo.', style='List Bullet')
    doc.add_paragraph('Batching agresivo: batch_size=32.', style='List Bullet')
    doc.add_paragraph('Fallback automático a CPU si ocurre OOM (Out Of Memory).', style='List Bullet')
    doc.add_paragraph('Secuencias cortas: 50 puntos (vs 100).', style='List Bullet')
    
    # ========== 4.6 ANÁLISIS HISTÉRESIS ==========
    doc.add_heading('4.6 Análisis de Histéresis Fuerza-Velocidad', level=1)
    
    doc.add_heading('4.6.1 Metodología del Sensor Virtual', level=2)
    doc.add_paragraph('Se implementó un sensor virtual de fuerza que permite detectar y cuantificar histéresis F-v sin necesidad de dinamómetro:')
    
    p = doc.add_paragraph()
    p.add_run('Pipeline de procesamiento:').bold = True
    doc.add_paragraph('Filtrado pasa-altas Butterworth (0.5 Hz, 2º orden).', style='List Bullet')
    doc.add_paragraph('Integración trapezoidal para obtener velocidad.', style='List Bullet')
    doc.add_paragraph('Recentrado a media cero.', style='List Bullet')
    doc.add_paragraph('Síntesis de fuerza mediante modelos de fricción.', style='List Bullet')
    
    doc.add_heading('4.6.2 Modelos de Fricción Implementados', level=2)
    
    p = doc.add_paragraph()
    p.add_run('Modelo Dahl:').bold = True
    eq = doc.add_paragraph()
    eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_equation(eq, r'\dot{z} = \alpha v \left(1 - \frac{z}{F_c}\text{sign}(v)\right), \quad \hat{F} = \sigma_0 z + \sigma_1 v')
    
    p = doc.add_paragraph()
    p.add_run('Modelo LuGre:').bold = True
    eq = doc.add_paragraph()
    eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_equation(eq, r'\dot{z} = v - \frac{\sigma_0|v|}{g(v)}z, \quad g(v) = F_c + (F_s - F_c)e^{-(v/v_s)^2}')
    eq = doc.add_paragraph()
    eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_equation(eq, r'\hat{F} = \sigma_0 z + \sigma_1 \dot{z} + \sigma_2 v')
    
    p = doc.add_paragraph()
    p.add_run('Modelo Bouc-Wen:').bold = True
    eq = doc.add_paragraph()
    eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_equation(eq, r'\dot{z} = Av - \beta|v||z|^{n-1}z - \gamma v|z|^n, \quad \hat{F} = k_v v + \alpha z')
    
    doc.add_heading('4.6.3 Detección de Eventos Transitorios', level=2)
    doc.add_paragraph('Se implementó detección automática de "pelitos" (eventos transitorios) mediante:')
    doc.add_paragraph('Umbrales adaptativos sobre picos (2.5σ).', style='List Bullet')
    doc.add_paragraph('Distancia mínima entre eventos.', style='List Bullet')
    doc.add_paragraph('Extracción de micro-ventanas (±100 muestras).', style='List Bullet')
    
    p = doc.add_paragraph()
    p.add_run('Resultado: ').bold = True
    p.add_run('En eventos transitorios, los lazos F-v son pronunciados y multi-valuados. El modelo Dahl '
              'resulta ganador con frecuencia (pre-deslizamiento dominante), mientras que LuGre captura '
              'efectos Stribeck cuando la velocidad cruza por cero.')
    
    doc.add_heading('4.6.4 Método de Cruce por Cero', level=2)
    doc.add_paragraph('Se implementó el método de cruce por cero para cálculo preciso de fricción dinámica:')
    doc.add_paragraph('Zona de cruce: |x| < 10% del stroke.', style='List Bullet')
    doc.add_paragraph('Separar por dirección de velocidad:', style='List Bullet')
    doc.add_paragraph('    • F_forward: Media de fuerza en avance (v > 0).')
    doc.add_paragraph('    • F_backward: Media de fuerza en retroceso (v < 0).')
    doc.add_paragraph('Fricción dinámica: F_friction_zc = (F_forward - F_backward) / 2.', style='List Bullet')
    
    p = doc.add_paragraph()
    p.add_run('Ventaja: ').bold = True
    p.add_run('Más preciso que el método pico a pico para ciclos de histéresis asimétricos.')
    
    # ========== 4.7 VALIDACIÓN ==========
    doc.add_heading('4.7 Validación Experimental', level=1)
    
    doc.add_heading('4.7.1 Calidad de Datos', level=2)
    doc.add_paragraph('Se verificó la calidad de los datos experimentales:')
    
    create_table(doc,
        ['Métrica', 'Resultado'],
        [
            ['NaN/Inf', '0 en todos los canales'],
            ['Tiempos monótonos', '✓ Verificado'],
            ['Outliers (|z|>6)', '<1% de muestras'],
            ['Coherencia temporal', 'Verificada']
        ]
    )
    doc.add_paragraph()
    
    doc.add_heading('4.7.2 Validación F = ma', level=2)
    doc.add_paragraph('Se implementó validación en tiempo real de la relación fuerza-aceleración:')
    
    eq = doc.add_paragraph()
    eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_equation(eq, r'F_{\text{medida}} \approx m \cdot a_{\text{sensor}}')
    
    doc.add_paragraph('Métricas calculadas:')
    doc.add_paragraph('Error porcentual entre F medida y m × a.', style='List Bullet')
    doc.add_paragraph('Transmisibilidad entre acelerómetros.', style='List Bullet')
    doc.add_paragraph('Desfase temporal entre señales.', style='List Bullet')
    
    doc.add_heading('4.7.3 Estado de Calibración', level=2)
    
    create_table(doc,
        ['Componente', 'Estado', 'Observaciones'],
        [
            ['Acelerómetros PCB 352C33', '✓ Validado', 'Sensibilidad verificada'],
            ['Celda de carga DYMH-105', '⚠ Pendiente', 'Requiere masas patrón'],
            ['Sincronización canales', '⚠ En proceso', 'Lag detectado en algunas sesiones']
        ]
    )
    doc.add_paragraph()
    
    # ========== 4.8 SOFTWARE ==========
    doc.add_heading('4.8 Infraestructura de Software', level=1)
    
    doc.add_heading('4.8.1 Archivos Principales Desarrollados', level=2)
    
    create_table(doc,
        ['Archivo', 'Líneas', 'Descripción'],
        [
            ['caracterizacion_sensor_fuerza.py', '1,085', 'GUI caracterización con shaker'],
            ['train_kanpinn_efficient.py', '878', 'Entrenamiento KAN-PINN optimizado'],
            ['interfaz_DAQ_acel_fuerza.py', '~800', 'Adquisición tiempo real'],
            ['analisis_frf_sistema.py', '395', 'Análisis FRF offline'],
            ['articulo_histeresis_fv.tex', '116', 'Artículo LaTeX IFToMM'],
            ['reporte_parcial.tex', '300', 'Documentación PINN']
        ]
    )
    doc.add_paragraph()
    
    doc.add_heading('4.8.2 Documentación Generada', level=2)
    doc.add_paragraph('RESUMEN_EJECUTIVO.txt: Guía rápida para Colab.', style='List Bullet')
    doc.add_paragraph('MAPA_COMPLETO_VERIFICADO.md: Auditoría detallada del código.', style='List Bullet')
    doc.add_paragraph('README_COLAB_COMPLETO.md: Instrucciones paso a paso.', style='List Bullet')
    doc.add_paragraph('Múltiples reportes de validación de datos.', style='List Bullet')
    
    # ========== 4.9 PROGRESO ==========
    doc.add_heading('4.9 Resultados Esperados y Progreso', level=1)
    
    doc.add_heading('4.9.1 Alineación con Objetivos', level=2)
    
    create_table(doc,
        ['Objetivo Específico', 'Progreso', 'Estado'],
        [
            ['Redactar 2 artículos científicos', '1 presentado, 1 en preparación', '50%'],
            ['Analizar muestra de sistemas mecánicos', 'Sistema DAQ caracterizado', '70%'],
            ['Diseñar modelo matemático de histéresis', 'PINN + modelos clásicos implementados', '80%'],
            ['Implementar modelo predictivo', 'KAN-PINN funcional', '60%'],
            ['Simular bajo diversas condiciones', 'DOE completado', '75%'],
            ['Optimizar algoritmos de IA', 'Optimizaciones GPU implementadas', '70%'],
            ['Evaluar impacto del modelo', 'Métricas comparativas obtenidas', '65%']
        ]
    )
    doc.add_paragraph()
    
    doc.add_heading('4.9.2 Métricas de Avance', level=2)
    
    create_table(doc,
        ['Indicador', 'Valor Actual'],
        [
            ['R² mejor modelo PINN', '0.50'],
            ['R² mejor modelo CNN', '0.9999'],
            ['Mejora PINN vs base', '+47%'],
            ['Speedup GPU vs CPU', '50-100x'],
            ['Sesiones experimentales', '14+'],
            ['Archivos de código', '25+']
        ]
    )
    doc.add_paragraph()
    
    # ========== 4.10 CONCLUSIONES ==========
    doc.add_heading('4.10 Conclusiones Parciales', level=1)
    
    p = doc.add_paragraph(style='List Number')
    p.add_run('Validación del enfoque PINN: ').bold = True
    p.add_run('Se demostró que incorporar restricciones físicas mejora significativamente el rendimiento '
              'predictivo (R² de 0.33 a 0.50).')
    
    p = doc.add_paragraph(style='List Number')
    p.add_run('Superioridad de CNN para alta frecuencia: ').bold = True
    p.add_run('Los modelos CNN superan dramáticamente a los PINN clásicos para datos experimentales '
              'de alta frecuencia (R² = 0.97 vs 0.03).')
    
    p = doc.add_paragraph(style='List Number')
    p.add_run('Identificación de parámetros: ').bold = True
    p.add_run('Se identificó un desafío clave: más grados de libertad no garantizan mejor rendimiento '
              'y pueden conducir a mínimos locales.')
    
    p = doc.add_paragraph(style='List Number')
    p.add_run('Sensor virtual de fuerza: ').bold = True
    p.add_run('Se validó la metodología para detectar histéresis F-v sin dinamómetro, usando modelos '
              'de fricción con memoria.')
    
    p = doc.add_paragraph(style='List Number')
    p.add_run('Infraestructura robusta: ').bold = True
    p.add_run('Se desarrolló un sistema completo de adquisición, procesamiento y análisis de datos experimentales.')
    
    # ========== 4.11 TRABAJO FUTURO ==========
    doc.add_heading('4.11 Trabajo Futuro', level=1)
    
    p = doc.add_paragraph()
    p.add_run('Corto plazo (próximo mes):').bold = True
    doc.add_paragraph('Completar calibración de celdas de carga con masas patrón.', style='List Bullet')
    doc.add_paragraph('Finalizar correcciones del artículo IFToMM.', style='List Bullet')
    doc.add_paragraph('Intensificar recolección de datos experimentales.', style='List Bullet')
    
    p = doc.add_paragraph()
    p.add_run('Mediano plazo (próximo semestre):').bold = True
    doc.add_paragraph('Desarrollar CNN híbridos que incorporen conocimiento físico.', style='List Bullet')
    doc.add_paragraph('Implementar incertidumbre bayesiana en predicciones.', style='List Bullet')
    doc.add_paragraph('Validar modelo en condiciones de corte reales.', style='List Bullet')
    
    p = doc.add_paragraph()
    p.add_run('Largo plazo (conclusión de tesis):').bold = True
    doc.add_paragraph('Integrar modelo predictivo en sistema de control.', style='List Bullet')
    doc.add_paragraph('Publicar segundo artículo en revista indexada.', style='List Bullet')
    doc.add_paragraph('Documentar metodología completa para reproducibilidad.', style='List Bullet')
    
    # ========== GUARDAR ==========
    output_path = r'c:\Users\jesus\Documents\Doctorado\Experimentos\CRio DAQ\cDAQ_9174\Capitulo4_Resultados.docx'
    doc.save(output_path)
    print(f'✓ Documento guardado en: {output_path}')
    return output_path

if __name__ == '__main__':
    main()
